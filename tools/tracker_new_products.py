"""Add the "New Products" 85% completion tab to CM_Open_Order_Tracker.xlsx.

Usage:
    python tools/tracker_new_products.py <in.xlsx> <out.xlsx>

What it adds (idempotent: an existing New Products tab is rebuilt):
  * Config!LaunchCompleteThreshold (defined name LaunchThreshold, default 0.85)
  * "New Products" tab (table tblNewProducts): one row per launch SKU.
    Inputs (yellow): SKU, Launch Name, Launch Date, Launch Qty, Source / Notes.
    Formulas pull every Open Orders line whose Item ends with the SKU
    (handles NetSuite "Parent : Child" item names) and compute
    Completed = Shipped/Received + FG at Factory, Basis = Launch Qty or
    Ordered, Completed %, and the 85% status / gap.
  * Summary block "NEW PRODUCT LAUNCHES" with the counts.
  * Read Me + Change Log rows.
  * Re-adds the Ethan Status dropdown on Ethan Updates (openpyxl drops the
    x14 extension form of that validation on save).

The seed SKU list is built from the workbook itself: SKUs that the
Supply Chain Update Meeting Workbook's "Launch Status 2026" sheet names
(confirmed) plus 2026 first-generation model codes (assumed; confirm).
"""
from __future__ import annotations

import datetime as dt
import sys
import warnings
from copy import copy

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo

warnings.filterwarnings("ignore")

SHEET = "New Products"
TABLE = "tblNewProducts"
THRESHOLD = 0.85
TODAY = dt.date(2026, 10, 16)
OO = "'Open Orders'!"
R = "$5:$%s$5000"  # row span used for every Open Orders reference


def oo(col: str) -> str:
    return f"{OO}${col}{R % col}"


# ---- which SKUs count as new products -------------------------------------
# Confirmed: literal SKUs / launch names that appear in the meeting workbook's
# "Launch Status 2026" sheet (see discovery/FORMULA_MAP.md).
CONFIRMED = {
    "P3 MAP TEAL": ("P3 Maple Teal", None, None),
    "SE PAN 10-": ("Panthera 10", dt.date(2026, 9, 15), None),
    "AIR RUSH RSE GLD": ("Air Rush Rose Gold", None, None),
}
PANTHERA_QTY = {"BCP PRE SE PAN 10-1 NW": 140, "BCP PRE SE PAN 10-2 LEA": 140, "BCP PRE SE PAN 10-3 WEB": 60}
# Assumed: product lines whose code is a 2026 first-generation model number.
ASSUMED_LINES = {
    "IKON5": "Ikon5",
    "BLA5": "Blak 5",
    "SPO3": "SPO3",
    "CND01": "Poison CND01",
    "DRA01": "Poison DRA01 (Dragon series)",
    "MAE01": "Poison MAE01",
    "NIT01": "Poison NIT01",
    "VX01": "Poison VX01",
}
CONF_SRC = "Confirmed — named on Launch Status 2026 (Supply Chain Update Meeting Workbook)"
ASSUMED_SRC = "Assumed — 2026 first-gen model code; confirm against Launch Status 2026"

HEADERS = [
    "SKU", "Launch Name", "Launch Date", "Launch Qty", "Vendor", "Product Type", "Product Line",
    "Lines", "Ordered", "Shipped / Received", "FG at Factory", "Completed", "In Production / Scheduled",
    "Open Qty", "Basis Qty", "Completed %", "85% Completion Status", "Units to Reach 85%",
    "Next Due Date", "Worst Days Late", "Days to Launch", "Source / Notes",
]
WIDTHS = [34, 26, 12, 10, 18, 13, 12, 7, 9, 11, 10, 10, 12, 9, 10, 11, 24, 11, 12, 10, 10, 60]
INPUT_COLS = {1, 2, 3, 4, 22}
INPUT_FILL = PatternFill("solid", fgColor="FFFFF2CC")
NUM = "#,##0;\\(#,##0\\);\\-"
DATE = "mm-dd-yy"


def seed_rows(wb):
    ws = wb["Open Orders"]
    hdr = [c.value for c in ws[4]]
    ix = {h: i for i, h in enumerate(hdr)}
    seen = {}
    for r in ws.iter_rows(min_row=5, values_only=True):
        if not r[0]:
            continue
        item = str(r[ix["Item"]] or "").strip()
        sku = item.split(" : ")[-1].strip()
        line = str(r[ix["Product Line"]] or "")
        if sku in seen:
            continue
        for key, (name, ldate, _) in CONFIRMED.items():
            if key in sku:
                seen[sku] = (name, ldate, PANTHERA_QTY.get(sku), CONF_SRC)
                break
        else:
            if line in ASSUMED_LINES:
                seen[sku] = (ASSUMED_LINES[line], None, None, ASSUMED_SRC)
    order = {n: i for i, n in enumerate([v[0] for v in CONFIRMED.values()] + list(ASSUMED_LINES.values()))}
    return sorted(seen.items(), key=lambda kv: (order.get(kv[1][0], 99), kv[0]))


def style_from(src, dst):
    dst.font = copy(src.font)
    dst.fill = copy(src.fill)
    dst.alignment = copy(src.alignment)
    dst.border = copy(src.border)
    dst.number_format = src.number_format


def formulas(r: int) -> dict[int, str]:
    crit = f"{oo('I')},\"*\"&$A{r}"
    sumif = lambda col: f"=SUMIFS({oo(col)},{crit})"
    status = (
        f'=IF(H{r}=0,"No order lines",IF(P{r}="","No basis qty",'
        f'IF(P{r}>=LaunchThreshold,"At or above "&TEXT(LaunchThreshold,"0%")&" complete",'
        f'"Below "&TEXT(LaunchThreshold,"0%")&" complete")))'
    )
    nxt = (
        f"_xlfn.MINIFS({oo('AE')},{crit},{oo('AE')},\">=\"&TODAY(),{oo('L')},\">0\")"
    )
    worst = f"_xlfn.MAXIFS({oo('AF')},{crit})"
    lookup = lambda col: f"=IFERROR(INDEX({oo(col)},MATCH(\"*\"&$A{r},{oo('I')},0)),\"\")"
    return {
        5: lookup("C"), 6: lookup("D"), 7: lookup("F"),
        8: f"=COUNTIFS({crit})",
        9: sumif("J"), 10: sumif("K"), 11: sumif("O"),
        12: f"=J{r}+K{r}",
        13: sumif("M"), 14: sumif("L"),
        15: f"=IF(D{r}>0,D{r},I{r})",
        16: f'=IF(O{r}>0,L{r}/O{r},"")',
        17: status,
        18: f'=IF(OR(H{r}=0,O{r}=0),"",MAX(0,ROUNDUP(O{r}*LaunchThreshold,0)-L{r}))',
        19: f'=IF({nxt}=0,"",{nxt})',
        20: f'=IF({worst}=0,"",{worst})',
        21: f'=IF(C{r}="","",C{r}-TODAY())',
    }


def build(src: str, dst: str) -> None:
    wb = openpyxl.load_workbook(src)
    nl = wb["New Lines"]  # style donor

    # -- Ethan Status dropdown (x14 extension is dropped by openpyxl) --------
    eu = wb["Ethan Updates"]
    if not any(dv.type == "list" for dv in eu.data_validations.dataValidation):
        dv = DataValidation(type="list", formula1="Lists!$A$2:$A$8", allow_blank=True)
        dv.add("F5:F1852")
        eu.add_data_validation(dv)

    # -- Config threshold ---------------------------------------------------
    cfg = wb["Config"]
    row = None
    for r in range(5, cfg.max_row + 1):
        if cfg.cell(r, 1).value == "LaunchCompleteThreshold":
            row = r
    if row is None:
        row = cfg.max_row + 1
        for c in range(1, 4):
            style_from(cfg.cell(row - 1, c), cfg.cell(row, c))
        cfg.cell(row, 1, "LaunchCompleteThreshold")
        cfg.cell(row, 2, THRESHOLD)
        cfg.cell(row, 3, "New Products tab: a launch SKU counts as complete when (Shipped/Received + FG at Factory) ÷ Basis Qty is at or above this. Basis Qty = Launch Qty if entered, else Ordered.")
        cfg.tables["tblConfig"].ref = f"A4:C{row}"
    cfg.cell(row, 2).number_format = "0%"
    wb.defined_names["LaunchThreshold"] = DefinedName("LaunchThreshold", attr_text=f"Config!$B${row}")

    # -- New Products sheet -------------------------------------------------
    if SHEET in wb.sheetnames:
        del wb[SHEET]
    ws = wb.create_sheet(SHEET, index=wb.sheetnames.index("New Lines") + 1)
    ws.sheet_properties.tabColor = "C8922A"
    ws["A1"] = "New Products — 85% completion status by launch SKU"
    style_from(nl["A1"], ws["A1"])
    ws["A2"] = ("Completed = Shipped/Received + FG at Factory across every Open Orders line whose Item ends with the SKU. "
                "Basis Qty = Launch Qty (yellow, from Launch Status 2026) if entered, else Ordered. "
                "Status uses Config › LaunchCompleteThreshold (85%). Yellow cells are inputs; add a SKU on the row under the table and the formulas fill down.")
    style_from(nl["A2"], ws["A2"])
    ws["A3"] = "Seed list: Confirmed rows are named on Launch Status 2026 in the Supply Chain Update Meeting Workbook; Assumed rows are 2026 first-generation model codes — delete any that are not launches, add the ones missing."
    style_from(nl["A2"], ws["A3"])
    ws.merge_cells("A1:K1")
    ws.merge_cells("A2:V2")
    ws.merge_cells("A3:V3")
    ws["A2"].alignment = Alignment(wrap_text=True, vertical="top")
    ws["A3"].alignment = Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[2].height = 30
    ws.row_dimensions[3].height = 28

    for c, h in enumerate(HEADERS, start=1):
        cell = ws.cell(4, c, h)
        style_from(nl["A4"], cell)
        ws.column_dimensions[get_column_letter(c)].width = WIDTHS[c - 1]
    ws.row_dimensions[4].height = 30

    body_font = Font(name="Arial", size=10)
    seeds = seed_rows(wb)
    r = 5
    for sku, (name, ldate, qty, srcnote) in seeds:
        ws.cell(r, 1, sku)
        ws.cell(r, 2, name)
        if ldate:
            ws.cell(r, 3, ldate)
        if qty:
            ws.cell(r, 4, qty)
        ws.cell(r, 22, srcnote)
        for c, f in formulas(r).items():
            ws.cell(r, c, f)
        for c in range(1, len(HEADERS) + 1):
            cell = ws.cell(r, c)
            cell.font = body_font
            if c in INPUT_COLS:
                cell.fill = INPUT_FILL
            if c in (3, 19):
                cell.number_format = DATE
            elif c == 16:
                cell.number_format = "0%"
            elif c in (4, 8, 9, 10, 11, 12, 13, 14, 15, 18, 20, 21):
                cell.number_format = NUM
            if c == 22:
                cell.alignment = Alignment(wrap_text=True, vertical="top")
        r += 1
    last = r - 1
    tbl = Table(displayName=TABLE, ref=f"A4:V{last}")
    donor = nl.tables["tblNewLines"].tableStyleInfo
    tbl.tableStyleInfo = TableStyleInfo(name=donor.name, showFirstColumn=donor.showFirstColumn,
                                        showLastColumn=donor.showLastColumn, showRowStripes=donor.showRowStripes,
                                        showColumnStripes=donor.showColumnStripes)
    ws.add_table(tbl)
    ws.freeze_panes = "B5"

    # -- Summary block --------------------------------------------------------
    sm = wb["Summary"]
    r0 = 47
    for rr in range(r0, r0 + 12):
        for c in range(1, 12):
            sm.cell(rr, c).value = None
    sm.cell(r0, 1, "NEW PRODUCT LAUNCHES — 85% COMPLETION")
    style_from(sm["A5"], sm.cell(r0, 1))
    sm.merge_cells(start_row=r0, start_column=1, end_row=r0, end_column=11)
    for c, h in enumerate(["Measure", "Value", "", "", "", "", "", "", "", "", "Note"], start=1):
        cell = sm.cell(r0 + 1, c, h or None)
        style_from(sm["A6"], cell)
    sm.merge_cells(start_row=r0 + 1, start_column=3, end_row=r0 + 1, end_column=10)
    Q = f"'{SHEET}'!$Q$5:$Q$500"
    rows = [
        ("Launch SKUs tracked", f"=COUNTA('{SHEET}'!$A$5:$A$500)", NUM, "Rows on the New Products tab"),
        ("At or above 85% complete", f'=COUNTIF({Q},"At or above*")', NUM, "Shipped/received + FG at factory ≥ 85% of basis qty"),
        ("Below 85% complete", f'=COUNTIF({Q},"Below*")', NUM, "Still short of the 85% launch target"),
        ("No order lines in tracker", f'=COUNTIF({Q},"No order lines")', NUM, "SKU on the list but not in any current source — check PO / SKU spelling"),
        ("Units completed vs basis (all launch SKUs)", f"=IF(SUM('{SHEET}'!$O$5:$O$500)=0,\"\",SUM('{SHEET}'!$L$5:$L$500)/SUM('{SHEET}'!$O$5:$O$500))", "0%", "Weighted by units"),
        ("Units still needed to reach 85%", f"=SUM('{SHEET}'!$R$5:$R$500)", NUM, "Sum of each SKU's gap to its 85% target"),
    ]
    for i, (label, f, nf, note) in enumerate(rows):
        rr = r0 + 2 + i
        a = sm.cell(rr, 1, label); style_from(sm["A7"], a); a.font = Font(name="Arial", size=10, bold=False)
        b = sm.cell(rr, 2, f); style_from(sm["B7"], b); b.number_format = nf
        k = sm.cell(rr, 11, note); style_from(sm["A43"], k); k.alignment = Alignment(wrap_text=False)
    note_r = r0 + 2 + len(rows)
    sm.cell(note_r, 1, "Completed = shipped/received + FG at factory. Basis = Launch Qty where entered on the New Products tab, else ordered qty. Threshold = Config › LaunchCompleteThreshold. Full SKU list: New Products tab.")
    style_from(sm["A43"], sm.cell(note_r, 1))
    sm.merge_cells(start_row=note_r, start_column=1, end_row=note_r, end_column=11)

    # -- Read Me ----------------------------------------------------------------
    rm = wb["Read Me"]
    rr = rm.max_row + 1
    rm.cell(rr, 1, "New products"); style_from(rm["A12"], rm.cell(rr, 1))
    rm.cell(rr, 2, "New Products = one row per 2026 launch SKU with its 85% completion status. Completed units = shipped/received + FG at factory from Open Orders; basis = Launch Qty (enter from Launch Status 2026) or ordered qty. Threshold lives on Config (LaunchCompleteThreshold). The HTML viewer shows the same on its New Products tab.")
    style_from(rm["B12"], rm.cell(rr, 2))

    # -- Change Log ---------------------------------------------------------------
    cl = wb["Change Log"]
    rr = cl.max_row + 1
    vals = [TODAY, "CB", "Structure",
            "Added New Products tab: 85% completion status per 2026 launch SKU (completed = shipped/received + FG at factory ÷ launch qty or ordered qty; threshold on Config). Summary block + HTML viewer 'New Products' view.",
            "Chase: reconcile the seed SKU list with Launch Status 2026 (meeting workbook); enter Launch Date / Launch Qty where known. Weekly: nothing extra — formulas recalc from Open Orders."]
    for c, v in enumerate(vals, start=1):
        cell = cl.cell(rr, c, v); style_from(cl.cell(rr - 1, c), cell)
    cl.tables["tblLog"].ref = f"A1:E{rr}"

    wb.save(dst)
    print(f"wrote {dst}: {len(seeds)} launch SKUs on '{SHEET}', threshold cell Config!B{row}")


if __name__ == "__main__":
    build(sys.argv[1], sys.argv[2])
